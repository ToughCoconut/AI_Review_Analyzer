import openai as op
from openpyxl import Workbook
from openpyxl import load_workbook

op.api_key = "Enter_Key_Here"

file_path = "Enter_Filepath_Here"

#ChatGPT Bot
def chat_with_chatgpt(prompt, model="text-davinci-003"):
    response = op.Completion.create(
        engine=model,
        prompt=prompt,

        max_tokens=1024,
        n=1,
        stop=None,
        temperature=0.5,
    )

    message = response.choices[0].text.strip()

    return message


def open_workbook(path):
    workbook_r = load_workbook(filename=path)

    #Can Work With Various Sheets In The Excel File
    sheet_r = workbook_r['Sheet_Name']
    print(sheet_r)
    print(f"The title of the Worksheet is: {sheet_r.title}")

    row = sheet_r.max_row
    column = sheet_r.max_column
  
    print("Total Rows:", row)
    print("Total Columns:", column)
    
    input_counter = 0
    cell_counter = 0

    total_GPT_responses_1 = ""
    total_GPT_responses_2 = ""
    for i in range(4, row + 1): 
        cell_obj = sheet_r.cell(row = i, column = 2) 
        #print(cell_obj.value)
        #print("\n")
        prompt1 = f"'{cell_obj.value}' \n Please act as a professional researcher to precisely list the key points for the following review in English, regarding the good, bad, and benefits, in point form and within 100 words"
        chatbot_response1 = chat_with_chatgpt(prompt1)
        #print(chatbot_response)
        sheet_r[f'C{i}'].value = chatbot_response1
        total_GPT_responses_1 = total_GPT_responses_1 + chatbot_response1 + "\n"
        input_counter += 1

        cell_counter += 1
        print(f"Cell number {cell_counter}")
        
        if (input_counter < 25) is False:
            print("Inside 2nd Prompt Collection")
            prompt2 = f"'{total_GPT_responses_1}' \n Please provide me with a list detailing the negatives, positives, and benefits of utilizing Facebook in point form and within 100 words. Additionally, kindly provide a concise summary of the collected reviews. The summary should highlight the most negative comments as well as the highest rating scores. Please remember that the objective of this research is to develop a value proposition for a digital business card networking platform based on product descriptions and user reviews. Note: The key points and summary only can be based on the information provided in the review."
            chatbot_response2 = chat_with_chatgpt(prompt2)
            print(chatbot_response2)
            total_GPT_responses_2 = total_GPT_responses_2 + chatbot_response2 + "\n"
            input_counter = 0
            total_GPT_responses_1 = ""
        
        if (cell_counter == row - 3) and (input_counter < 25):
            print("Inside 2nd Prompt Collection")
            prompt2 = f"'{total_GPT_responses_1}' \n Please provide me with a list detailing the negatives, positives, and benefits of utilizing Facebook in point form and within 100 words. Additionally, kindly provide a concise summary of the collected reviews. The summary should highlight the most negative comments as well as the highest rating scores. Please remember that the objective of this research is to develop a value proposition for a digital business card networking platform based on product descriptions and user reviews. Note: The key points and summary only can be based on the information provided in the review."
            chatbot_response2 = chat_with_chatgpt(prompt2)
            print(chatbot_response2)
            total_GPT_responses_2 = total_GPT_responses_2 + chatbot_response2 + "\n"
            input_counter = 0
            total_GPT_responses_1 = ""
    
    print("Going to summarize the 2nd prompt collection")
    prompt3 = f"'{total_GPT_responses_2}' \n Analyze the aspects of these inputs and give a response in the exact same structure with the same characteristics. The summary should highlight the most negative comments as well as the highest rating scores. Please remember that the objective of this research is to develop a value proposition for a digital business card networking platform based on product descriptions and user reviews. Note: The key points and summary only can be based on the information provided in the review."
    chatbot_response3 = chat_with_chatgpt(prompt3)
    print(chatbot_response3)
    sheet_r['D4'].value = chatbot_response3


    workbook_r.save(path)
    print("Done!")
        


if __name__ == "__main__":
    open_workbook(file_path)
